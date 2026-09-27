"""Builds .docx deliverables with python-docx."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..db import models


def _add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_impact_onepager(project: models.Project, analysis: models.ChangeImpactAnalysisRow) -> Document:
    doc = Document()

    title = doc.add_heading(f"Change Impact Assessment", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(project.name).runs[0].bold = True
    doc.add_paragraph(analysis.change_title)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(analysis.change_summary)

    doc.add_heading("Stakeholder Impact", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, label in enumerate(["Stakeholder", "Impact Severity", "A-D-K-A-R", "Barrier", "Heat", "What Changes"]):
        hdr[i].text = label

    for s in analysis.stakeholder_impacts:
        row = table.add_row().cells
        row[0].text = s.stakeholder
        row[1].text = f"{s.impact_severity}/5"
        row[2].text = (
            f"{s.adkar_awareness}-{s.adkar_desire}-{s.adkar_knowledge}-"
            f"{s.adkar_ability}-{s.adkar_reinforcement}"
        )
        row[3].text = s.adkar_barrier_dimension or "None"
        row[4].text = s.heat_rating
        row[5].text = s.what_changes

    doc.add_heading("Overall Risks", level=1)
    _add_bullets(doc, analysis.overall_risks)

    doc.add_heading("Success Factors", level=1)
    _add_bullets(doc, analysis.success_factors)

    package = project.communication_package
    if package:
        doc.add_heading("Change Readiness Summary", level=1)
        p = doc.add_paragraph()
        run = p.add_run(f"Readiness score: {package.readiness_score}/100")
        run.bold = True
        run.font.size = Pt(13)
        doc.add_paragraph(package.readiness_rationale)
        doc.add_heading("Top Risks", level=2)
        _add_bullets(doc, package.readiness_top_risks)
        doc.add_heading("Recommended Next Steps", level=2)
        _add_bullets(doc, package.readiness_next_steps)

    return doc


def build_comms_package(
    project: models.Project,
    package: models.CommunicationPackageRow,
    training_checklist: list[models.ChecklistItemRow],
    implementation_checklist: list[models.ChecklistItemRow],
) -> Document:
    doc = Document()
    doc.add_heading("Change Communication Package", level=0)
    doc.add_paragraph(project.name).runs[0].bold = True

    doc.add_heading("Employee Announcement", level=1)
    doc.add_heading(package.employee_subject, level=2)
    for para in package.employee_body.split("\n\n"):
        doc.add_paragraph(para)

    doc.add_heading("Manager Communication", level=1)
    doc.add_heading(package.manager_subject, level=2)
    for para in package.manager_body.split("\n\n"):
        doc.add_paragraph(para)
    doc.add_heading("Talking Points", level=2)
    _add_bullets(doc, package.manager_talking_points)

    doc.add_heading("FAQ", level=1)
    for item in package.faq:
        p = doc.add_paragraph()
        p.add_run(f"Q: {item['question']}").bold = True
        doc.add_paragraph(f"A: {item['answer']}")

    def checklist_table(heading: str, items: list[models.ChecklistItemRow]) -> None:
        doc.add_heading(heading, level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, label in enumerate(["Item", "Owner", "Due", "Done"]):
            hdr[i].text = label
        for it in items:
            row = table.add_row().cells
            row[0].text = it.item
            row[1].text = it.owner
            row[2].text = it.due or "-"
            row[3].text = "Yes" if it.is_done else "No"

    checklist_table("Training Checklist", training_checklist)
    checklist_table("Implementation Checklist", implementation_checklist)

    return doc
