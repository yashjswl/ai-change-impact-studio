"""Builds a 5-slide executive summary .pptx with python-pptx. Heat map cells
are colored natively via table cell fills, no image/matplotlib dependency,
so the deck stays fully editable in PowerPoint."""

from __future__ import annotations

import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ..db import models

HEAT_COLORS = {
    "Red": RGBColor(0xC0, 0x39, 0x2B),
    "Amber": RGBColor(0xE1, 0xA1, 0x00),
    "Green": RGBColor(0x27, 0x82, 0x3B),
}
HEADER_COLOR = RGBColor(0x1F, 0x2A, 0x44)


def _add_title_slide(prs: Presentation, project: models.Project) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = project.name
    subtitle = slide.placeholders[1]
    subtitle.text = (
        f"Change Impact & Readiness, Executive Summary\n"
        f"Prepared {datetime.date.today().isoformat()}"
    )


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str], empty_message: str) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    items = bullets or [empty_message]
    body.text = items[0]
    for item in items[1:]:
        p = body.add_paragraph()
        p.text = item


def _add_heatmap_slide(prs: Presentation, stakeholders: list[models.StakeholderImpactRow]) -> None:
    layout = prs.slide_layouts[5]  # title only
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Stakeholder Impact vs. Readiness Heat Map"

    rows = len(stakeholders) + 1
    cols = 4
    left, top, width, height = Inches(0.5), Inches(1.5), Inches(9), Inches(0.4 * rows)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    headers = ["Stakeholder", "Impact", "Readiness", "Rating"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_COLOR
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(12)

    for r, s in enumerate(stakeholders, start=1):
        table.cell(r, 0).text = s.stakeholder
        table.cell(r, 1).text = f"{s.impact_score:.0f}"
        table.cell(r, 2).text = f"{s.readiness_score:.0f}"
        rating_cell = table.cell(r, 3)
        rating_cell.text = s.heat_rating
        rating_cell.fill.solid()
        rating_cell.fill.fore_color.rgb = HEAT_COLORS.get(s.heat_rating, RGBColor(0x80, 0x80, 0x80))
        for p in rating_cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True

    if not stakeholders:
        table.cell(1, 0).text = "No impact analysis generated yet"


def _add_raci_slide(prs: Presentation, raci_items: list[models.RaciItemRow]) -> None:
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "RACI Summary (Top Activities)"

    top_items = raci_items[:8]
    rows = len(top_items) + 1
    cols = 3
    left, top, width, height = Inches(0.5), Inches(1.5), Inches(9), Inches(0.4 * rows)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    headers = ["Activity", "Responsible", "Accountable"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_COLOR
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r, item in enumerate(top_items, start=1):
        table.cell(r, 0).text = item.activity
        table.cell(r, 1).text = item.responsible
        table.cell(r, 2).text = item.accountable

    if not top_items:
        table.cell(1, 0).text = "No RACI matrix generated yet"


def build_executive_summary(
    project: models.Project,
    stakeholders: list[models.StakeholderImpactRow],
    top_risks: list[str],
    raci_items: list[models.RaciItemRow],
    next_steps: list[str],
) -> Presentation:
    prs = Presentation()
    _add_title_slide(prs, project)
    _add_heatmap_slide(prs, stakeholders)
    _add_bullet_slide(prs, "Top Risks", top_risks, "No high/critical risks logged yet")
    _add_raci_slide(prs, raci_items)
    _add_bullet_slide(prs, "Recommended Next Steps", next_steps, "No next steps generated yet")
    return prs
